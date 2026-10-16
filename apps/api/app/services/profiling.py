"""Deterministic CSV/XLSX parsing, schema inference and data quality
profiling. No LLM involved anywhere in this module — every number it
produces is a reproducible calculation over the actual uploaded data.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

import polars as pl
from charset_normalizer import from_bytes
from openpyxl import load_workbook

MAX_SAMPLE_VALUES = 5
MAX_PROFILED_ROWS = 200_000  # profile a bounded sample for very large files
_CANDIDATE_DELIMITERS = [",", ";", "\t", "|"]

_DATE_PATTERNS = [
    re.compile(r"^\d{4}-\d{2}-\d{2}$"),  # ISO: 2026-03-05
    re.compile(r"^\d{2}/\d{2}/\d{4}$"),  # 05/03/2026 (DD/MM or MM/DD, ambiguous)
    re.compile(r"^\d{2}-\d{2}-\d{4}$"),  # 05-03-2026
]
_NUMERIC_TEXT_PATTERN = re.compile(r"^[-+]?[£$€]?\s*[\d,]+(\.\d+)?%?$")


class ParsingError(Exception):
    pass


@dataclass
class ColumnProfile:
    name: str
    inferred_type: str  # integer | float | boolean | date | datetime | text
    suggested_type: str  # inferred_type, or a better guess when flagged below
    null_count: int
    null_ratio: float
    distinct_count: int
    sample_values: list[Any]
    stats: dict[str, Any] = field(default_factory=dict)
    flags: list[str] = field(default_factory=list)  # e.g. "numeric_stored_as_text"


@dataclass
class DataQualityIssue:
    severity: str  # "info" | "warning"
    issue_type: str
    column: str | None
    detail: str


@dataclass
class TableProfile:
    sheet_name: str
    row_count: int
    column_count: int
    duplicate_row_count: int
    columns: list[ColumnProfile]
    issues: list[DataQualityIssue]
    quality_score: float


def _decode_csv_bytes(content: bytes) -> str:
    best = from_bytes(content).best()
    if best is None:
        # Last-resort fallback: most real-world "unknown encoding" CSVs
        # from UK business systems are Windows-1252/Latin-1.
        return content.decode("latin-1", errors="replace")
    return str(best)


def _sniff_delimiter(sample: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="".join(_CANDIDATE_DELIMITERS))
        return dialect.delimiter
    except csv.Error:
        # Fall back to whichever candidate appears most often in the header line.
        header_line = sample.splitlines()[0] if sample.splitlines() else ""
        counts = {d: header_line.count(d) for d in _CANDIDATE_DELIMITERS}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else ","


def parse_csv(content: bytes) -> dict[str, pl.DataFrame]:
    text = _decode_csv_bytes(content)
    delimiter = _sniff_delimiter(text[:8192])

    try:
        # try_parse_dates is deliberately NOT used here: Polars infers a
        # single date format per column from a sample, then silently nulls
        # out any value that doesn't match it — which would quietly lose
        # real data in a column with mixed date formats (common in
        # real-world exports). Date columns are instead detected via the
        # text-pattern heuristic in _detect_text_flags, below, which
        # degrades to "text" rather than losing values.
        df = pl.read_csv(
            io.StringIO(text),
            separator=delimiter,
            infer_schema_length=10_000,
            ignore_errors=True,
            n_rows=MAX_PROFILED_ROWS,
        )
    except Exception as exc:
        raise ParsingError(f"Could not parse CSV file: {exc}") from exc

    if df.width == 0:
        raise ParsingError("CSV file has no columns.")

    return {"Sheet1": df}


def list_excel_sheets(content: bytes) -> list[str]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ParsingError(f"Could not open Excel workbook: {exc}") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def parse_excel(content: bytes) -> dict[str, pl.DataFrame]:
    sheet_names = list_excel_sheets(content)
    if not sheet_names:
        raise ParsingError("Workbook contains no worksheets.")

    tables: dict[str, pl.DataFrame] = {}
    for sheet_name in sheet_names:
        try:
            df = pl.read_excel(io.BytesIO(content), sheet_name=sheet_name)
        except Exception as exc:
            raise ParsingError(f"Could not parse sheet '{sheet_name}': {exc}") from exc
        if df.width > 0:
            tables[sheet_name] = df

    if not tables:
        raise ParsingError("No worksheet in the workbook contains any columns.")

    return tables


def parse_upload(extension: str, content: bytes) -> dict[str, pl.DataFrame]:
    if extension == ".csv":
        return parse_csv(content)
    if extension == ".xlsx":
        return parse_excel(content)
    raise ParsingError(f"Unsupported extension for parsing: {extension}")


def _polars_dtype_to_label(dtype: pl.DataType) -> str:
    if dtype in (pl.Int8, pl.Int16, pl.Int32, pl.Int64, pl.UInt8, pl.UInt16, pl.UInt32, pl.UInt64):
        return "integer"
    if dtype in (pl.Float32, pl.Float64):
        return "float"
    if dtype == pl.Boolean:
        return "boolean"
    if dtype == pl.Date:
        return "date"
    if dtype in (pl.Datetime, pl.Datetime("us")):
        return "datetime"
    return "text"


def _detect_text_flags(values: list[str]) -> tuple[str | None, list[str]]:
    """For a text column, check whether it looks like numbers or dates
    that were stored as text (a common real-world data quality issue)."""
    non_empty = [v for v in values if v and v.strip()]
    if not non_empty:
        return None, []

    date_matches = sum(1 for v in non_empty if any(p.match(v.strip()) for p in _DATE_PATTERNS))
    if date_matches / len(non_empty) >= 0.9:
        return "date", ["date_stored_as_text"]

    numeric_matches = sum(1 for v in non_empty if _NUMERIC_TEXT_PATTERN.match(v.strip()))
    if numeric_matches / len(non_empty) >= 0.9:
        return "float", ["numeric_stored_as_text"]

    return None, []


def _profile_column(series: pl.Series) -> ColumnProfile:
    name = series.name
    total = len(series)
    null_count = series.null_count()
    base_type = _polars_dtype_to_label(series.dtype)
    non_null = series.drop_nulls()
    distinct_count = non_null.n_unique()
    sample_values = non_null.unique().head(MAX_SAMPLE_VALUES).to_list()

    suggested_type = base_type
    flags: list[str] = []
    stats: dict[str, Any] = {}

    if base_type == "text":
        text_values = [str(v) for v in non_null.to_list()]
        detected_type, text_flags = _detect_text_flags(text_values)
        if detected_type:
            suggested_type = detected_type
            flags.extend(text_flags)
        lengths = [len(v) for v in text_values] if text_values else []
        if lengths:
            stats["min_length"] = min(lengths)
            stats["max_length"] = max(lengths)
    elif base_type in ("integer", "float"):
        if len(non_null) > 0:
            stats["min"] = non_null.min()
            stats["max"] = non_null.max()
            stats["mean"] = round(float(non_null.mean()), 4) if non_null.mean() is not None else None
            stats["median"] = non_null.median()
            stats["std_dev"] = round(float(non_null.std()), 4) if len(non_null) > 1 else 0.0
    elif base_type == "boolean":
        stats["true_count"] = int(non_null.sum()) if len(non_null) else 0
        stats["false_count"] = len(non_null) - stats["true_count"]
    elif base_type in ("date", "datetime"):
        if len(non_null) > 0:
            stats["min"] = str(non_null.min())
            stats["max"] = str(non_null.max())

    if total > 1 and distinct_count == 1 and null_count == 0:
        flags.append("constant_column")

    return ColumnProfile(
        name=name,
        inferred_type=base_type,
        suggested_type=suggested_type,
        null_count=null_count,
        null_ratio=round(null_count / total, 4) if total else 0.0,
        distinct_count=distinct_count,
        sample_values=[_jsonable(v) for v in sample_values],
        stats={k: _jsonable(v) for k, v in stats.items()},
        flags=flags,
    )


def _jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def _compute_quality_score(
    overall_null_ratio: float, duplicate_ratio: float, flagged_column_ratio: float
) -> float:
    """Deterministic 0-100 score. Not a statistical gold standard — a
    transparent, documented heuristic: missing data costs up to 40 points,
    duplicate rows up to 30, and columns with detected formatting issues
    (e.g. numbers stored as text) up to 20. See docs/ARCHITECTURE.md."""
    score = 100.0
    score -= min(40.0, overall_null_ratio * 100 * 0.4)
    score -= min(30.0, duplicate_ratio * 100 * 0.3)
    score -= min(20.0, flagged_column_ratio * 100 * 0.2)
    return round(max(0.0, min(100.0, score)), 1)


def profile_table(sheet_name: str, df: pl.DataFrame) -> TableProfile:
    row_count = df.height
    column_count = df.width
    duplicate_row_count = row_count - df.unique().height if row_count else 0

    columns = [_profile_column(df[col]) for col in df.columns]

    issues: list[DataQualityIssue] = []
    for col in columns:
        if col.null_ratio == 1.0:
            issues.append(
                DataQualityIssue("warning", "fully_empty_column", col.name, "Every value is missing.")
            )
        elif col.null_ratio >= 0.2:
            issues.append(
                DataQualityIssue(
                    "warning",
                    "missing_values",
                    col.name,
                    f"{col.null_ratio:.0%} of values are missing.",
                )
            )
        elif col.null_ratio > 0:
            issues.append(
                DataQualityIssue(
                    "info", "missing_values", col.name, f"{col.null_ratio:.0%} of values are missing."
                )
            )

        if "numeric_stored_as_text" in col.flags:
            issues.append(
                DataQualityIssue(
                    "warning",
                    "numeric_stored_as_text",
                    col.name,
                    "Values look numeric but are stored as text.",
                )
            )
        if "date_stored_as_text" in col.flags:
            issues.append(
                DataQualityIssue(
                    "warning", "date_stored_as_text", col.name, "Values look like dates but are stored as text."
                )
            )
        if "constant_column" in col.flags:
            issues.append(
                DataQualityIssue(
                    "info", "constant_column", col.name, "Every non-null value in this column is identical."
                )
            )

    if duplicate_row_count > 0:
        issues.append(
            DataQualityIssue(
                "warning",
                "duplicate_rows",
                None,
                f"{duplicate_row_count} duplicate row(s) found out of {row_count}.",
            )
        )

    overall_null_ratio = (
        sum(c.null_count for c in columns) / (row_count * column_count) if row_count and column_count else 0.0
    )
    duplicate_ratio = duplicate_row_count / row_count if row_count else 0.0
    flagged_column_ratio = (
        sum(1 for c in columns if c.flags) / column_count if column_count else 0.0
    )
    quality_score = _compute_quality_score(overall_null_ratio, duplicate_ratio, flagged_column_ratio)

    return TableProfile(
        sheet_name=sheet_name,
        row_count=row_count,
        column_count=column_count,
        duplicate_row_count=duplicate_row_count,
        columns=columns,
        issues=issues,
        quality_score=quality_score,
    )
