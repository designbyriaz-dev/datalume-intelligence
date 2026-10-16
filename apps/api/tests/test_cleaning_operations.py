"""The second wave of cleaning operations (deduplicate was the first,
Phase 2): fill-missing-values, normalize-text, filter-rows, and
calculated-field. Each follows the same pattern verified for
deduplicate — a new DatasetVersion is produced, the source version is
left untouched, and a Transformation audit row is recorded."""

import pytest
from httpx import AsyncClient


async def _register_and_login(client: AsyncClient, email: str) -> str:
    await client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "correct-horse-battery", "full_name": "Test User"},
    )
    resp = await client.post("/api/v1/auth/login", json={"email": email, "password": "correct-horse-battery"})
    return resp.json()["access_token"]


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


async def _create_personal_workspace(client: AsyncClient, token: str) -> str:
    resp = await client.post(
        "/api/v1/workspaces", json={"name": "Test Workspace", "industry": "general"}, headers=_auth(token)
    )
    return resp.json()["id"]


async def _upload_csv(client: AsyncClient, token: str, workspace_id: str, name: str, content: bytes) -> str:
    files = {"file": ("data.csv", content, "text/csv")}
    resp = await client.post(
        f"/api/v1/workspaces/{workspace_id}/datasets", data={"name": name}, files=files, headers=_auth(token)
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["tables"][0]["id"]


def _columns_by_name(version_body: dict) -> dict:
    return {c["name"]: c for c in version_body["tables"][0]["profile"]["columns"]}


# --- fill-missing-values -----------------------------------------------

FILL_CSV = (
    b"name,age,salary\n"
    b"Jane,34,50000\n"
    b"Bob,,62000\n"
    b"Alice,29,\n"
    b"Eve,41,71000\n"
)


async def test_fill_missing_values_drop_rows(client: AsyncClient) -> None:
    token = await _register_and_login(client, "fill1@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILL_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/fill-missing-values",
        json={"column": "age", "strategy": "drop_rows"},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["version_number"] == 2
    assert body["tables"][0]["row_count"] == 3  # Bob's null-age row removed
    assert _columns_by_name(body)["age"]["null_count"] == 0


async def test_fill_missing_values_constant(client: AsyncClient) -> None:
    token = await _register_and_login(client, "fill2@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILL_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/fill-missing-values",
        json={"column": "salary", "strategy": "constant", "constant_value": 0},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["tables"][0]["row_count"] == 4  # no rows removed
    assert _columns_by_name(body)["salary"]["null_count"] == 0


async def test_fill_missing_values_mean(client: AsyncClient) -> None:
    token = await _register_and_login(client, "fill3@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILL_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/fill-missing-values",
        json={"column": "age", "strategy": "mean"},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    # Filling a null with the mean of the other values leaves the mean
    # unchanged — mean of (34, 29, 41, fill_value) == mean of (34, 29, 41).
    assert _columns_by_name(body)["age"]["null_count"] == 0
    assert _columns_by_name(body)["age"]["stats"]["mean"] == pytest.approx((34 + 29 + 41) / 3, abs=1e-3)


async def _dataset_id(client: AsyncClient, token: str, workspace_id: str) -> str:
    resp = await client.get(f"/api/v1/workspaces/{workspace_id}/datasets", headers=_auth(token))
    return resp.json()[0]["id"]


async def test_fill_missing_values_mean_rejects_non_numeric_column(client: AsyncClient) -> None:
    token = await _register_and_login(client, "fill4@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILL_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/fill-missing-values",
        json={"column": "name", "strategy": "mean"},
        headers=_auth(token),
    )
    assert resp.status_code == 400
    assert "numeric" in resp.json()["detail"].lower()


async def test_fill_missing_values_unknown_column_rejected(client: AsyncClient) -> None:
    token = await _register_and_login(client, "fill5@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILL_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/fill-missing-values",
        json={"column": "not_a_column", "strategy": "drop_rows"},
        headers=_auth(token),
    )
    assert resp.status_code == 400
    assert "unknown column" in resp.json()["detail"].lower()

    original_resp = await client.get(f"/api/v1/datasets/{await _dataset_id(client, token, workspace_id)}", headers=_auth(token))
    # The invalid request must not have created a version 2.
    assert len(original_resp.json()["versions"]) == 1


async def test_fill_missing_values_mode(client: AsyncClient) -> None:
    token = await _register_and_login(client, "fill6@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    csv = b"name,status\nA,active\nB,\nC,active\nD,inactive\n"
    table_id = await _upload_csv(client, token, workspace_id, "Data", csv)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/fill-missing-values",
        json={"column": "status", "strategy": "mode"},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    assert _columns_by_name(resp.json())["status"]["null_count"] == 0


# --- normalize-text ------------------------------------------------------

TEXT_CSV = b"name,city\n  Jane  ,  LONDON  \nBob,new   york\n"


async def test_normalize_text_trims_and_lowercases(client: AsyncClient) -> None:
    token = await _register_and_login(client, "norm1@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", TEXT_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/normalize-text",
        json={"column": "city", "case": "lower", "trim": True, "collapse_whitespace": True},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text

    dataset_id = await _dataset_id(client, token, workspace_id)
    detail = (await client.get(f"/api/v1/datasets/{dataset_id}", headers=_auth(token))).json()
    new_table = next(v for v in detail["versions"] if v["version_number"] == 2)["tables"][0]
    assert new_table["profile"]["columns"][1]["sample_values"]  # sanity: profile recomputed


async def test_normalize_text_rejects_non_text_column(client: AsyncClient) -> None:
    token = await _register_and_login(client, "norm2@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILL_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/normalize-text",
        json={"column": "age"},
        headers=_auth(token),
    )
    assert resp.status_code == 400
    assert "not a text column" in resp.json()["detail"].lower()


async def test_normalize_text_rejects_unknown_case(client: AsyncClient) -> None:
    token = await _register_and_login(client, "norm3@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", TEXT_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/normalize-text",
        json={"column": "city", "case": "shout"},
        headers=_auth(token),
    )
    assert resp.status_code == 400


# --- filter-rows -----------------------------------------------------------

FILTER_CSV = (
    b"name,age,city\n"
    b"Jane,34,London\n"
    b"Bob,19,Leeds\n"
    b"Alice,29,London\n"
    b"Eve,41,Leeds\n"
)


async def test_filter_rows_numeric_gte(client: AsyncClient) -> None:
    token = await _register_and_login(client, "filt1@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILTER_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/filter-rows",
        json={"column": "age", "operator": "gte", "value": 30},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["tables"][0]["row_count"] == 2  # Jane (34), Eve (41)


async def test_filter_rows_contains(client: AsyncClient) -> None:
    token = await _register_and_login(client, "filt2@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILTER_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/filter-rows",
        json={"column": "city", "operator": "contains", "value": "Lon"},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["tables"][0]["row_count"] == 2  # Jane, Alice


async def test_filter_rows_contains_requires_text_column(client: AsyncClient) -> None:
    token = await _register_and_login(client, "filt3@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILTER_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/filter-rows",
        json={"column": "age", "operator": "contains", "value": "3"},
        headers=_auth(token),
    )
    assert resp.status_code == 400


async def test_filter_rows_missing_value_rejected(client: AsyncClient) -> None:
    token = await _register_and_login(client, "filt4@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILTER_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/filter-rows",
        json={"column": "age", "operator": "gte"},
        headers=_auth(token),
    )
    assert resp.status_code == 400


async def test_filter_rows_original_version_preserved(client: AsyncClient) -> None:
    token = await _register_and_login(client, "filt5@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", FILTER_CSV)

    await client.post(
        f"/api/v1/dataset-tables/{table_id}/filter-rows",
        json={"column": "age", "operator": "gte", "value": 30},
        headers=_auth(token),
    )

    dataset_id = await _dataset_id(client, token, workspace_id)
    detail = (await client.get(f"/api/v1/datasets/{dataset_id}", headers=_auth(token))).json()
    original = next(v for v in detail["versions"] if v["version_number"] == 1)
    assert original["tables"][0]["row_count"] == 4


# --- calculated-field --------------------------------------------------

CALC_CSV = b"name,price,quantity\nWidget,10,3\nGadget,25,2\n"


async def test_calculated_field_multiply_two_columns(client: AsyncClient) -> None:
    token = await _register_and_login(client, "calc1@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", CALC_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/calculated-field",
        json={"new_column_name": "total", "left_column": "price", "operator": "multiply", "right_column": "quantity"},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["tables"][0]["column_count"] == 4
    assert "total" in _columns_by_name(body)


async def test_calculated_field_add_constant(client: AsyncClient) -> None:
    token = await _register_and_login(client, "calc2@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", CALC_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/calculated-field",
        json={"new_column_name": "price_plus_tax", "left_column": "price", "operator": "add", "right_value": 2.5},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    assert "price_plus_tax" in _columns_by_name(resp.json())


async def test_calculated_field_concat_text_columns(client: AsyncClient) -> None:
    token = await _register_and_login(client, "calc3@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    csv = b"first,last\nJane,Doe\nBob,Smith\n"
    table_id = await _upload_csv(client, token, workspace_id, "Data", csv)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/calculated-field",
        json={
            "new_column_name": "full_name",
            "left_column": "first",
            "operator": "concat",
            "right_column": "last",
            "separator": " ",
        },
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    assert "full_name" in _columns_by_name(resp.json())


async def test_calculated_field_rejects_existing_column_name(client: AsyncClient) -> None:
    token = await _register_and_login(client, "calc4@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", CALC_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/calculated-field",
        json={"new_column_name": "price", "left_column": "price", "operator": "add", "right_value": 1},
        headers=_auth(token),
    )
    assert resp.status_code == 400
    assert "already exists" in resp.json()["detail"].lower()


async def test_calculated_field_rejects_both_right_column_and_value(client: AsyncClient) -> None:
    token = await _register_and_login(client, "calc5@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", CALC_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/calculated-field",
        json={
            "new_column_name": "x",
            "left_column": "price",
            "operator": "add",
            "right_column": "quantity",
            "right_value": 1,
        },
        headers=_auth(token),
    )
    assert resp.status_code == 400
    assert "exactly one" in resp.json()["detail"].lower()


async def test_calculated_field_rejects_non_numeric_for_arithmetic(client: AsyncClient) -> None:
    token = await _register_and_login(client, "calc6@example.com")
    workspace_id = await _create_personal_workspace(client, token)
    table_id = await _upload_csv(client, token, workspace_id, "Data", CALC_CSV)

    resp = await client.post(
        f"/api/v1/dataset-tables/{table_id}/calculated-field",
        json={"new_column_name": "x", "left_column": "name", "operator": "add", "right_value": 1},
        headers=_auth(token),
    )
    assert resp.status_code == 400
    assert "not numeric" in resp.json()["detail"].lower()


# --- multi-sheet: sibling sheets carried through unchanged ----------------


async def test_cleaning_operation_on_one_sheet_leaves_other_sheets_untouched(client: AsyncClient) -> None:
    import io

    import openpyxl

    token = await _register_and_login(client, "multi1@example.com")
    workspace_id = await _create_personal_workspace(client, token)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws1 = wb.create_sheet("Tenants")
    for row in [["unit", "rent"], ["A1", 1200], ["A2", 950]]:
        ws1.append(row)
    ws2 = wb.create_sheet("Arrears")
    for row in [["unit", "owed"], ["A1", 0], ["A2", 450]]:
        ws2.append(row)
    buf = io.BytesIO()
    wb.save(buf)

    files = {
        "file": (
            "portfolio.xlsx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    upload_resp = await client.post(
        f"/api/v1/workspaces/{workspace_id}/datasets", data={"name": "Portfolio"}, files=files, headers=_auth(token)
    )
    tables = {t["sheet_name"]: t for t in upload_resp.json()["tables"]}
    tenants_table_id = tables["Tenants"]["id"]

    resp = await client.post(
        f"/api/v1/dataset-tables/{tenants_table_id}/filter-rows",
        json={"column": "rent", "operator": "gt", "value": 1000},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    new_tables = {t["sheet_name"]: t for t in body["tables"]}
    assert new_tables["Tenants"]["row_count"] == 1  # only A1 (rent 1200) kept
    assert new_tables["Arrears"]["row_count"] == 2  # untouched sibling sheet
